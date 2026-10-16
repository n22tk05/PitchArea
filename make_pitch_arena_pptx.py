import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Colors
    BG_COLOR = RGBColor(10, 14, 23)
    PRIMARY_BLUE = RGBColor(37, 99, 235)
    LIGHT_BLUE = RGBColor(96, 165, 250)
    CYAN = RGBColor(6, 182, 212)
    WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(148, 163, 184)
    CARD_BG = RGBColor(16, 24, 40)
    CARD_BORDER = RGBColor(35, 45, 65)
    AMBER = RGBColor(245, 158, 11)
    ROSE = RGBColor(244, 63, 94)
    EMERALD = RGBColor(16, 185, 129)

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_header(slide, badge_text, title_text, subtitle_text, slide_num_str):
        # Header box
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10.5), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = f"● {badge_text.upper()}"
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = CYAN

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p1.space_before = Pt(4)

        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(3)

        # Slide number box
        num_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.4), Inches(1.2), Inches(0.5))
        ntf = num_box.text_frame
        np = ntf.paragraphs[0]
        np.alignment = PP_ALIGN.RIGHT
        np.text = slide_num_str
        np.font.size = Pt(14)
        np.font.bold = True
        np.font.color.rgb = LIGHT_BLUE

    # ------------------ SLIDE 1: TITLE ------------------
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1)

    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
    ttf = tbox.text_frame
    ttf.word_wrap = True
    ttf.margin_left = ttf.margin_top = ttf.margin_right = ttf.margin_bottom = 0

    p_badge = ttf.paragraphs[0]
    p_badge.alignment = PP_ALIGN.CENTER
    p_badge.text = "HỆ THỐNG HUẤN LUYỆN & MÔ PHỎNG PHẢN BIỆN KHỞI NGHIỆP THỜI GIAN THỰC"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = CYAN

    p_title = ttf.add_paragraph()
    p_title.alignment = PP_ALIGN.CENTER
    p_title.text = "PITCH·ARENA STUDIO"
    p_title.font.size = Pt(46)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.space_before = Pt(14)

    p_sub = ttf.add_paragraph()
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.text = "The Gamified AI Defense Arena with In-Memory RAG & Multi-Agent Deliberation (MAD)\nĐấu trường phản biện AI đối thoại giọng nói ~220ms, thẩm định khách quan và xuất phụ lục sửa Slide"
    p_sub.font.size = Pt(15)
    p_sub.font.color.rgb = TEXT_MUTED
    p_sub.space_before = Pt(14)

    p_foot = ttf.add_paragraph()
    p_foot.alignment = PP_ALIGN.CENTER
    p_foot.text = "In-Memory RAG • Realtime Voice STT • FSM Arena Engine • MAD Arbitration • Actionable Appendix\nPRODUCT REVIEW & TECHNICAL SHOWCASE DECK • 2026"
    p_foot.font.size = Pt(12)
    p_foot.font.color.rgb = LIGHT_BLUE
    p_foot.space_before = Pt(28)

    # ------------------ SLIDE 2: PAIN POINT ------------------
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2)
    add_header(s2, "Bối Cảnh & Thực Trạng", "Tại Sao 85% Đội Thi Thất Bại Trong Vòng Phản Biện Q&A?", "Các đội thi SV-Startup và gọi vốn thất bại không phải vì ý tưởng tồi mà vì thiếu môi trường rèn luyện phản xạ thực tế.", "02 / 07")

    stats = [
        ("85%", "Bị 'Hạ Gục' Ở 3 Phút Q&A", "Lúng túng khi Hội đồng xoáy sâu vào số liệu và tính khả thi"),
        ("30 Giây", "Áp Lực Phản Xạ Cực Độ", "Dễ rơi vào bẫy ngụy biện, nói vòng vo hoặc im lặng quá 2 giây"),
        ("0 Giải Pháp", "Huấn Luyện Tự Động Trước Đây", "Chỉ nhờ thầy cô nghe thử thủ công; ChatGPT không có áp lực thời gian")
    ]
    for i, (num, lbl, sub) in enumerate(stats):
        left = Inches(0.8 + i * 4.0)
        shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(3.7), Inches(1.4))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = num
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = LIGHT_BLUE
        p2 = tf.add_paragraph()
        p2.text = lbl
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p3 = tf.add_paragraph()
        p3.text = sub
        p3.font.size = Pt(9)
        p3.font.color.rgb = TEXT_MUTED

    # 2 Comparison cards
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.7), Inches(5.7), Inches(3.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "❌ HẠN CHẾ KHI LUYỆN TẬP THỦ CÔNG"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ROSE
    points1 = [
        "Phản hồi chủ quan: Người nghe thử ngại va chạm, không bóc trần đúng điểm mù.",
        "Không có áp lực thời gian: Không có đồng hồ đếm ngược 30s hay phát hiện ngập ngừng.",
        "Không có giải pháp sửa đổi: Chỉ chê 'chưa thuyết phục' mà không biết sửa cụ thể slide nào."
    ]
    for pt in points1:
        p = tf1.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(6)

    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.7), Inches(5.7), Inches(3.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = PRIMARY_BLUE
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "💡 GIẢI PHÁP PITCH·ARENA ĐỘT PHÁ"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CYAN
    points2 = [
        "Đấu trường Voice 220ms: Tương tác giọng nói thời gian thực với VAD bắt ngập ngừng.",
        "Hội đồng 4 Giám khảo: Xoáy từ Thị trường, Công nghệ, Unit Economics đến Rủi ro sao chép.",
        "Báo cáo Phụ lục hành động: Tự động chỉ rõ lỗi ngụy biện và gợi ý slide phụ lục đối kháng."
    ]
    for pt in points2:
        p = tf2.add_paragraph()
        p.text = f"✔ {pt}"
        p.font.size = Pt(10)
        p.font.color.rgb = WHITE
        p.space_before = Pt(6)

    # ------------------ SLIDE 3: 4 WEAPONS ------------------
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3)
    add_header(s3, "Kiến Trúc Kỹ Thuật", "4 Vũ Khí Kỹ Thuật Cốt Lõi Tạo Nên PITCH·ARENA", "Kết hợp đồng bộ giữa RAG trong bộ nhớ, Realtime Voice, cỗ máy FSM và Hội đồng MAD Chamber.", "03 / 07")

    weapons = [
        ("1. In-Memory RAG", "BÓC TÁCH & WHITELIST", "Bóc tách file Word .docx thành 5 khối đề mục, trích xuất số liệu tiền tệ, phần trăm vào bảng Whitelist đối soát tức thời.", CYAN),
        ("2. Realtime Voice", "220MS STT & VAD", "Nhận diện giọng nói STT trực tiếp, tích hợp VAD Radar bắt khoảng lặng 1.5s - 3.0s để Boss Giám khảo dồn ép thí sinh.", EMERALD),
        ("3. FSM Combat State", "CỖ MÁY TRẠNG THÁI", "Đệm suy nghĩ 7s, 30s đối chất, Tạm dừng chiến thuật, Đóng băng thời gian và Khóa sàn tân thủ 20% HP.", AMBER),
        ("4. MAD Chamber", "PHÂN XỬ ĐA TÁC TỬ", "Hội đồng đa tác tử: Công tố viên (bắt bẻ), Luật sư (bào chữa), Trọng tài Arbiter tính điểm Q = 0.4D + 0.35F + 0.25H.", ROSE)
    ]
    for i, (title, sub, desc, color) in enumerate(weapons):
        left = Inches(0.8 + i * 3.0)
        shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(2.8), Inches(4.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(9)
        p2.font.bold = True
        p2.font.color.rgb = color
        p2.space_before = Pt(3)
        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(12)

    # ------------------ SLIDE 4: 4 JURY BOSSES ------------------
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4)
    add_header(s4, "Hội Đồng Giám Khảo", "Hội Đồng 4 Trụ Cột Thẩm Định Chuyên Môn Khắc Nghiệt", "Mỗi Giám khảo mang một cá tính độc lập, chất vấn đúng nhóm đề mục tài liệu được phân công.", "04 / 07")

    bosses = [
        ("Shark Trần Nam", "LV.98 • THỊ TRƯỜNG", "Vấn đề & Thị trường", "Bóc trần khảo sát hình thức, nhu cầu ảo và khả năng chi tiền của người dùng.", ROSE),
        ("TS. Lê Minh Trang", "LV.92 • CÔNG NGHỆ", "Sản phẩm & Tech", "Truy vấn kiến trúc kỹ thuật, độ trễ, tự phát triển thuật toán hay mượn API ngoài.", CYAN),
        ("GS. Vũ Hoàng", "LV.95 • TÀI CHÍNH", "Unit Economics", "Lạnh lùng bóc trần từng con số, chi phí làm mẫu thử, CAC, LTV và điểm hòa vốn.", AMBER),
        ("ThS. Đặng Mai Lan", "LV.94 • RỦI RO", "Moat & Lộ trình", "Chất vấn rào cản phòng thủ sao chép, nguy cơ tan rã đội ngũ và tính khả thi 12 tháng.", LIGHT_BLUE)
    ]
    for i, (name, title, domain, desc, color) in enumerate(bosses):
        left = Inches(0.8 + i * 3.0)
        shape = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(2.8), Inches(4.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = name
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(9)
        p2.font.bold = True
        p2.font.color.rgb = color
        p2.space_before = Pt(2)
        p3 = tf.add_paragraph()
        p3.text = f"Lĩnh vực: {domain}"
        p3.font.size = Pt(10)
        p3.font.color.rgb = LIGHT_BLUE
        p3.space_before = Pt(6)
        p4 = tf.add_paragraph()
        p4.text = desc
        p4.font.size = Pt(10)
        p4.font.color.rgb = TEXT_MUTED
        p4.space_before = Pt(8)

    # ------------------ SLIDE 5: 5 STEPS WORKFLOW ------------------
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5)
    add_header(s5, "Quy Trình Hoạt Động", "5 Bước Trải Nghiệm Thực Chiến Tại PITCH·ARENA", "Thiết kế mạch lạc từ lúc kéo thả tài liệu đến khi nhận báo cáo và bước vào vòng tái đấu.", "05 / 07")

    steps = [
        ("BƯỚC 1", "Nạp Tài Liệu", "Kéo thả đề án .docx. Hệ thống bóc tách 5 khối và nạp vào bộ nhớ RAG trong 3 giây."),
        ("BƯỚC 2", "Cấu Hình Sảnh", "Chọn Full Arena hoặc Solo Boss, chọn độ khó (Easy/Normal/Hardcore) và Preset thẩm định."),
        ("BƯỚC 3", "Thuyết Minh", "Thuyết minh đề tài qua mic. Hệ thống hiển thị sóng âm và đếm tốc độ WPM trực tiếp."),
        ("BƯỚC 4", "Phản Biện", "Boss truyền câu hỏi realtime. Đệm 7s suy nghĩ. 30s đối chất gay cấn. Trọng tài phân xử HP."),
        ("BƯỚC 5", "Báo Cáo Phụ Lục", "Xem biểu đồ biến thiên HP, bảng phụ lục 3 cột hành động và bấm Re-match luyện lại.")
    ]
    for i, (step_num, title, desc) in enumerate(steps):
        left = Inches(0.8 + i * 2.4)
        shape = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(2.25), Inches(4.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = step_num
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = CYAN
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(3)
        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(10)

    # ------------------ SLIDE 6: ACTIONABLE APPENDIX ------------------
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6)
    add_header(s6, "Đầu Ra Cốt Lõi", "Báo Cáo Phụ Lục Hành Động: Biến Thất Bại Thành Vũ Khí", "Tự động chỉ rõ lỗi ngụy biện và xuất bảng hướng dẫn chi tiết slide phụ lục cần bổ sung.", "06 / 07")

    cols = [
        ("Cột 1: Câu Hỏi & Bối Cảnh", "Ghi lại chính xác câu hỏi Boss chất vấn và câu trả lời của thí sinh qua micro, đối soát tài liệu Word.", CYAN),
        ("Cột 2: Bóc Trần Ngụy Biện", "Chỉ rõ điểm yếu chí tử: Nói vòng vo, Đánh tráo khái niệm, Ảo tưởng dữ liệu mẫu hoặc thiếu kiểm chứng.", AMBER),
        ("Cột 3: Slide Hành Động", "Đề xuất cụ thể tiêu đề Slide phụ lục cần tạo thêm (A1, A2...) kèm luận cứ bảo vệ mẫu để đưa ngay vào bài.", EMERALD)
    ]
    for i, (title, desc, color) in enumerate(cols):
        left = Inches(0.8 + i * 4.0)
        shape = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(3.7), Inches(3.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = color
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(8)

    # Rematch box
    rm_shape = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.4))
    rm_shape.fill.solid()
    rm_shape.fill.fore_color.rgb = CARD_BG
    rm_shape.line.color.rgb = PRIMARY_BLUE
    rm_tf = rm_shape.text_frame
    rm_tf.word_wrap = True
    p = rm_tf.paragraphs[0]
    p.text = "📈 VÒNG LẶP HUẤN LUYỆN TINH GỌN (QUICK RE-MATCH LOOP)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = LIGHT_BLUE
    p2 = rm_tf.add_paragraph()
    p2.text = "Giữ nguyên tài liệu RAG, tự động ghi nhận câu hỏi đã hỏi và lọc bỏ các câu có độ tương đồng ≥ 80%, giúp thí sinh rèn luyện nhiều góc chất vấn mới mà không cần nạp lại file."
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(3)

    # ------------------ SLIDE 7: SUMMARY ------------------
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s7)
    add_header(s7, "Tổng Kết & Giá Trị", "PITCH·ARENA: Vũ Khí Chiến Lược Cho Mọi Nhà Khởi Nghiệp", "Nâng tầm bản lĩnh phản biện của sinh viên và người trẻ Việt Nam trước Hội đồng giám khảo.", "07 / 07")

    summary_stats = [
        ("100%", "Bám Sát Đề Án Thật", "Không suy đoán hoang đường, bám sát từng bảng biểu trong file Word"),
        ("300%", "Tăng Tốc Phản Xạ", "Rèn luyện thói quen tư duy súc tích, trung thực dưới áp lực 30 giây"),
        ("1 Click", "Sẵn Sàng Triển Khai", "Trải nghiệm ngay trên trình duyệt, không cần cài đặt phần mềm phức tạp")
    ]
    for i, (num, lbl, sub) in enumerate(summary_stats):
        left = Inches(0.8 + i * 4.0)
        shape = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(3.7), Inches(1.4))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = num
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = LIGHT_BLUE
        p2 = tf.add_paragraph()
        p2.text = lbl
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p3 = tf.add_paragraph()
        p3.text = sub
        p3.font.size = Pt(9)
        p3.font.color.rgb = TEXT_MUTED

    # Bottom summary box
    sum_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.7), Inches(11.7), Inches(3.1))
    sum_box.fill.solid()
    sum_box.fill.fore_color.rgb = CARD_BG
    sum_box.line.color.rgb = PRIMARY_BLUE
    stf = sum_box.text_frame
    stf.word_wrap = True
    p = stf.paragraphs[0]
    p.text = "🎯 ĐÁNH GIÁ SẢN PHẨM & GIÁ TRỊ THỰC TIỄN (PRODUCT REVIEW VERDICT)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CYAN
    bullets = [
        "Độ hoàn thiện: Hoàn thành đầy đủ 5 Phase từ In-Memory RAG, WebSocket Voice 220ms, FSM Engine đến MAD Chamber và Appendix Report.",
        "Độ tin cậy: Toàn bộ codebase TypeScript đồng bộ 100% không lỗi (Zero Type Errors), kiểm chứng nghiêm ngặt qua test cases.",
        "Dân chủ hóa cố vấn: Mọi nhóm sinh viên ở mọi miền đều được tiếp cận Hội đồng Giám khảo tiêu chuẩn quốc tế 24/7 hoàn toàn miễn phí.",
        "★ PITCH·ARENA — BƯỚC ĐẦU CỦA MỌI CHIẾN THẮNG KHỞI NGHIỆP ★"
    ]
    for b in bullets:
        p = stf.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(10)
        p.font.color.rgb = WHITE
        p.space_before = Pt(5)

    out_path = r"C:\Users\ADMIN\Desktop\slide_pitch_arena.pptx"
    prs.save(out_path)
    print(f"Successfully generated {out_path}")

if __name__ == "__main__":
    create_deck()
